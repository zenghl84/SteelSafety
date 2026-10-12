"""Read the application workbook without changing its values or source file."""
from datetime import datetime
from pathlib import Path
from math import isfinite
from decimal import Decimal
from statistics import median

from openpyxl import load_workbook

WORKBOOK_NAME = "商洛煤棚监测_应用数据 (1).xlsx"
POINT_IDS = tuple(f"S{i}" for i in range(1, 11))


def source_path(project_dir: Path) -> Path:
    external = project_dir.parent / WORKBOOK_NAME
    return external if external.is_file() else project_dir / "data" / WORKBOOK_NAME


def number(value):
    if value is None or str(value).strip() == "":
        return None
    result = float(value)
    return result if isfinite(result) else None


def select_daily_median_record(records: list[dict]) -> dict:
    """Select one measured row on the latest day, without averaging its fields."""
    latest_day = max(row["timestamp"].date() for row in records)
    candidates = [
        row for row in records
        if row["timestamp"].date() == latest_day
        and row["corrected_strain"] is not None
        and isfinite(row["corrected_strain"])
    ]
    if not candidates:
        return {"date": latest_day, "record": None, "median": None, "sample_count": 0}
    # Decimal keeps equally distant measurements tied at the source precision.
    middle = median(Decimal(str(row["corrected_strain"])) for row in candidates)
    distances = [abs(Decimal(str(row["corrected_strain"])) - middle) for row in candidates]
    closest = min(distances)
    selected = max(
        (row for row, distance in zip(candidates, distances) if distance == closest),
        key=lambda row: row["timestamp"],
    )
    return {
        "date": latest_day, "record": selected,
        "median": float(middle), "sample_count": len(candidates),
    }


def load_monitoring_data(path: Path) -> dict:
    records_by_point = {}
    count = 0
    workbook = load_workbook(path, read_only=True, data_only=True)
    required = {"时间", "传感器编号", "传感器型号", "测点名称", "测点位置", "应变变化值", "单位1", "温度", "单位2"}
    try:
        for sheet in workbook:
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, ())
            if sheet.title not in POINT_IDS:
                continue
            if not required.issubset(headers):
                raise ValueError(f"工作表 {sheet.title} 缺少必要的数据列")
            for row_number, values in enumerate(rows, 2):
                if not any(value is not None for value in values):
                    continue
                row = dict(zip(headers, values))
                point = str(row["测点名称"]).strip()
                if point not in POINT_IDS:
                    raise ValueError(f"{sheet.title} 第 {row_number} 行测点名称无效：{point}")
                timestamp = row["时间"]
                if not isinstance(timestamp, datetime):
                    timestamp = datetime.fromisoformat(str(timestamp))
                record = {
                    "point": point,
                    "timestamp": timestamp,
                    "location": str(row["测点位置"] or "未提供位置"),
                    "sensor_id": str(row["传感器编号"] or "未提供"),
                    "sensor_model": str(row["传感器型号"] or "未提供"),
                    "strain": number(row["应变变化值"]),
                    "strain_unit": str(row["单位1"] or ""),
                    "corrected_strain": number(row.get("温度修正后应变")),
                    "corrected_strain_unit": str(row.get("单位4") or row["单位1"] or ""),
                    "temperature": number(row["温度"]),
                    "temperature_unit": str(row["单位2"] or ""),
                }
                count += 1
                records_by_point.setdefault(point, []).append(record)
    finally:
        workbook.close()
    if not records_by_point:
        raise ValueError("应用数据表中没有可读取的测点记录")
    return {
        "points": {
            point: select_daily_median_record(records)
            for point, records in records_by_point.items()
        },
        "record_count": count,
    }


def format_value(value):
    return "—" if value is None else f"{value:.4f}".rstrip("0").rstrip(".")
