"""Read monitoring data and apply the original workbook's temperature correction."""
from datetime import datetime, timedelta
from pathlib import Path
from math import isfinite
from decimal import Decimal
from statistics import median

from openpyxl import load_workbook

WORKBOOK_NAME = "商洛煤棚监测_应用数据.xlsx"
POINT_IDS = tuple(f"S{i}" for i in range(1, 11))
# Verified against all 3,870 formulas in the v1.2 monitoring workbook:
# corrected strain = G + 0.2 * (I - K), with the original sensor baselines.
TEMPERATURE_CORRECTION_COEFFICIENT = Decimal("0.2")
POINT_CALIBRATIONS = {
    "S1": ("0226026496", 22.5),
    "S2": ("0226026497", 23.1),
    "S3": ("0226026474", 22.6),
    "S5": ("0226026498", 22.2),
    "S6": ("0226026471", 22.8),
    "S7": ("0226026495", 22.2),
    "S8": ("0226026475", 22.0),
    "S9": ("0226026473", 22.6),
    "S10": ("0226026494", 22.5),
}


def source_path(project_dir: Path) -> Path:
    bundled = project_dir / "data" / WORKBOOK_NAME
    if bundled.is_file():
        return bundled
    external = project_dir.parent / WORKBOOK_NAME
    return external if external.is_file() else bundled


def number(value):
    if value is None or str(value).strip() == "":
        return None
    result = float(value)
    return result if isfinite(result) else None


def calculate_corrected_strain(strain, temperature, reference_temperature):
    """Use the original +0.2 temperature term, preserving missing measurements."""
    values = [number(value) for value in (strain, temperature, reference_temperature)]
    if any(value is None for value in values):
        return None
    strain, temperature, reference_temperature = map(lambda value: Decimal(str(value)), values)
    return float(
        strain + TEMPERATURE_CORRECTION_COEFFICIENT * (temperature - reference_temperature)
    )


def select_daily_median_record(records: list[dict]) -> dict:
    """Select one measured row on the latest day, without averaging its fields."""
    latest_day = max(row["timestamp"].date() for row in records)
    candidates = [
        row for row in records
        if row["timestamp"].date() == latest_day
        and row["corrected_strain"] is not None
        and isfinite(row["corrected_strain"])
    ]
    if not candidates:
        return {"date": latest_day, "record": None, "median": None, "sample_count": 0}
    # Decimal keeps equally distant measurements tied at the source precision.
    middle = median(Decimal(str(row["corrected_strain"])) for row in candidates)
    distances = [abs(Decimal(str(row["corrected_strain"])) - middle) for row in candidates]
    closest = min(distances)
    selected = max(
        (row for row, distance in zip(candidates, distances) if distance == closest),
        key=lambda row: row["timestamp"],
    )
    return {
        "date": latest_day, "record": selected,
        "median": float(middle), "sample_count": len(candidates),
    }


def load_monitoring_data(path: Path) -> dict:
    records_by_point = {}
    count = 0
    calculated_count = 0
    workbook = load_workbook(path, read_only=True, data_only=True)
    required = {"时间", "传感器编号", "传感器型号", "测点名称", "测点位置", "应变变化值", "单位1", "温度", "单位2"}
    try:
        for sheet in workbook:
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, ())
            if sheet.title not in POINT_IDS:
                continue
            if not required.issubset(headers):
                raise ValueError(f"工作表 {sheet.title} 缺少必要的数据列")
            for row_number, values in enumerate(rows, 2):
                if not any(value is not None for value in values):
                    continue
                row = dict(zip(headers, values))
                point = str(row["测点名称"]).strip()
                if point not in POINT_IDS:
                    raise ValueError(f"{sheet.title} 第 {row_number} 行测点名称无效：{point}")
                timestamp = row["时间"]
                if not isinstance(timestamp, datetime):
                    timestamp = datetime.fromisoformat(str(timestamp))
                record = {
                    "point": point,
                    "timestamp": timestamp,
                    "location": str(row["测点位置"] or "未提供位置"),
                    "sensor_id": str(row["传感器编号"] or "未提供"),
                    "sensor_model": str(row["传感器型号"] or "未提供"),
                    "strain": number(row["应变变化值"]),
                    "strain_unit": str(row["单位1"] or ""),
                    "corrected_strain": number(row.get("温度修正后应变")),
                    "corrected_strain_unit": str(row.get("单位4") or row["单位1"] or ""),
                    "temperature": number(row["温度"]),
                    "temperature_unit": str(row["单位2"] or ""),
                }
                record["reference_temperature"] = number(row.get("基准温度"))
                record["correction_source"] = (
                    "workbook" if record["corrected_strain"] is not None else None
                )
                if record["corrected_strain"] is None:
                    if record["strain"] is not None and record["temperature"] is not None:
                        if record["reference_temperature"] is None:
                            calibration = POINT_CALIBRATIONS.get(point)
                            if calibration is None or calibration[0] != record["sensor_id"]:
                                raise ValueError(
                                    f"{sheet.title} 第 {row_number} 行传感器无匹配的原基准温度，"
                                    "无法计算温度修正后应变"
                                )
                            record["reference_temperature"] = calibration[1]
                        record["corrected_strain"] = calculate_corrected_strain(
                            record["strain"], record["temperature"], record["reference_temperature"]
                        )
                        record["correction_source"] = "calculated"
                        calculated_count += 1
                count += 1
                records_by_point.setdefault(point, []).append(record)
    finally:
        workbook.close()
    if not records_by_point:
        raise ValueError("应用数据表中没有可读取的测点记录")
    return {
        "points": {
            point: select_daily_median_record(records)
            for point, records in records_by_point.items()
        },
        "record_count": count,
        "calculated_record_count": calculated_count,
        "history": {
            point: sorted(records, key=lambda row: row["timestamp"])
            for point, records in records_by_point.items()
        },
    }


def select_recent_records(records: list[dict], days: int = 3) -> list[dict]:
    """Use the point's latest recorded day and preceding calendar days."""
    if days < 1:
        raise ValueError("监测天数必须至少为一天")
    if not records:
        return []
    latest = max(row["timestamp"] for row in records)
    start = datetime.combine(
        latest.date() - timedelta(days=days - 1), datetime.min.time()
    )
    return sorted(
        (row for row in records if start <= row["timestamp"] <= latest),
        key=lambda row: row["timestamp"],
    )


def format_value(value):
    return "—" if value is None else f"{value:.4f}".rstrip("0").rstrip(".")
